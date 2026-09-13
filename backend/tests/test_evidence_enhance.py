"""Muse-branch seams: live-only ingestion endpoints, verbatim proposal checks."""
from app.jobs import brief_substantive
from app.models import CompanyBrief, EvidenceItem
from app.storage import Store
from app.extraction import apply_proposal, propose_evidence, normalized


def test_live_ingestion_endpoints_require_live(client):
    assert client.post('/api/evidence/enhance').status_code == 503
    assert client.post('/api/sources/nosuchid/detect-target').status_code == 503


def test_source_delete_cascades_evidence_and_bumps_version(client):
    client.post('/api/sessions', json={'demo': False})
    upload = client.post('/api/sources', data={'slot': 'resume', 'text': 'Experience\nBuilt APIs with Python.'})
    assert upload.status_code == 201
    before = client.get('/api/evidence').json()
    assert len(before['evidence']) > 0
    sid = upload.json()['source_id']
    deleted = client.delete('/api/sources/' + sid)
    assert deleted.status_code == 200
    after = client.get('/api/evidence').json()
    assert after['sources'] == [] and after['evidence'] == []
    assert after['session']['input_version'] == before['session']['input_version'] + 1
    assert client.delete('/api/sources/' + sid).status_code == 404


def test_brief_substantive_thresholds():
    assert brief_substantive(CompanyBrief(summary='x' * 100, key_facts=['y' * 20] * 3))
    assert not brief_substantive(CompanyBrief(summary='short', key_facts=['y' * 20] * 3))
    assert not brief_substantive(CompanyBrief(summary='x' * 100, key_facts=['too short', 'y' * 20]))
    assert not brief_substantive(CompanyBrief(summary='x' * 100, key_facts=[]))


def _session_store(tmp_path):
    store = Store(tmp_path / 'd')
    session = store.add('session', 'a' * 32, {'input_version': 0})
    return store, session


def test_apply_proposal_keeps_only_verbatim(tmp_path):
    store, session = _session_store(tmp_path)
    source = {'id': 's1', 'locator_map': [{'locator': 'paragraph 1', 'text': 'Built payment APIs using Python. Led a team of four.'}]}
    items = [EvidenceItem(locator='paragraph 1', exact_excerpt='Built payment APIs using Python.', category='Experience'),
             EvidenceItem(locator='paragraph 1', exact_excerpt='Invented quantum revenue.', category='Experience'),
             EvidenceItem(locator='paragraph 99', exact_excerpt='Built payment APIs using Python.', category='Experience')]
    accepted, dropped = apply_proposal(store, session, source, items)
    assert (accepted, dropped) == (1, 2)
    evidence = store.list('evidence', session['id'])
    assert len(evidence) == 1 and evidence[0]['category'] == 'Experience'
    assert evidence[0]['support_status'] == 'self_reported'


def test_static_merge_skips_covered(tmp_path):
    store, session = _session_store(tmp_path)
    source = {'id': 's1', 'locator_map': [{'locator': 'line 1', 'text': 'Built payment APIs using Python.'},
                                          {'locator': 'line 2', 'text': 'Won hackathon prize in 2024.'}]}
    propose_evidence(store, session, source, skip_covered=normalized('Built payment APIs using Python.'))
    texts = [e['exact_excerpt'] for e in store.list('evidence', session['id'])]
    assert texts == ['Won hackathon prize in 2024.']


def test_docx_heading_styles_and_labels_are_not_evidence(tmp_path):
    import io
    from docx import Document
    from app.extraction import extract
    doc = Document()
    doc.add_heading('Experience', level=2)
    doc.add_paragraph('Led a launch growing revenue by 10%.')
    doc.add_paragraph('Skills')
    doc.add_paragraph('Python, SQL.')
    buf = io.BytesIO()
    doc.save(buf)
    chunks = extract(buf.getvalue(), 'r.docx')
    store, session = _session_store(tmp_path)
    propose_evidence(store, session, {'id': 's1', 'locator_map': chunks})
    got = [(e['category'], e['exact_excerpt']) for e in store.list('evidence', session['id'])]
    assert got == [('Experience', 'Led a launch growing revenue by 10%.'),
                   ('Skills', 'Python, SQL.')]


def test_summary_heading_maps_to_summary(tmp_path):
    from app.extraction import extract
    chunks = extract(b'# Jane Doe\n\n## Summary\n\nPM with 3 years of fintech experience.\n', 'r.md')
    store, session = _session_store(tmp_path)
    propose_evidence(store, session, {'id': 's1', 'locator_map': chunks})
    got = [(e['category'], e['exact_excerpt']) for e in store.list('evidence', session['id'])]
    assert got == [('Summary', 'PM with 3 years of fintech experience.')]


def test_detect_header_extracts_name_and_contact_without_model(client):
    client.post('/api/sessions', json={'demo': False})
    upload = client.post('/api/sources', data={'slot': 'resume', 'text': 'DHRUV NAWANI\nProduct Manager\nreach me at dhruv@example.com or 8140789789\n'})
    assert upload.status_code == 201
    found = client.post('/api/sources/' + upload.json()['source_id'] + '/detect-header')
    assert found.status_code == 200
    assert found.json()['name'] == 'DHRUV NAWANI'
    assert found.json()['contact'] == 'dhruv@example.com | 8140789789'


def test_detect_header_skips_dates_and_labels(client):
    client.post('/api/sessions', json={'demo': False})
    upload = client.post('/api/sources', data={'slot': 'resume', 'text': '## Experience\n05/2026 - Present\nSKILLS\nPython, SQL.'})
    assert upload.status_code == 201
    found = client.post('/api/sources/' + upload.json()['source_id'] + '/detect-header').json()
    assert found['name'] == ''
    assert found['contact'] == ''


def test_detect_header_rejects_non_candidate_source(client):
    client.post('/api/sessions', json={'demo': False})
    upload = client.post('/api/sources', data={'slot': 'job', 'text': 'Hiring engineers.'})
    assert upload.status_code == 201
    assert client.post('/api/sources/' + upload.json()['source_id'] + '/detect-header').status_code == 404
